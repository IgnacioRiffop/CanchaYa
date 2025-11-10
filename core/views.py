from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.models import User
from django.contrib import messages
from django.contrib.auth.tokens import default_token_generator
from django.utils.http import urlsafe_base64_encode, urlsafe_base64_decode
from django.utils.encoding import force_bytes, force_str
from django.core.mail import send_mail
from django.contrib.auth.hashers import make_password
from django.db import transaction
from core.models import Usuario
from django.http import JsonResponse
import re
from .forms import *
from .models import *
from django.core.paginator import Paginator
from datetime import time
from datetime import date, timedelta
from datetime import datetime
import json
import stripe
from django.conf import settings
from django.db.models import Sum, Count
from django.http import JsonResponse
import io
import pandas as pd
from django.http import HttpResponse
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from .models import Reserva
from openpyxl.utils import get_column_letter
from openpyxl.styles import numbers
from django.db.models.functions import ExtractWeekDay
from django.http import HttpResponse
from openpyxl.utils import get_column_letter
from openpyxl.styles import numbers
from django.db.models import Count
from django.db.models.functions import ExtractWeekDay
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors





def contacto(request):
    return render(request, 'core/contacto.html')

def promociones(request):
    return render(request, 'core/promociones.html')

from django.shortcuts import render, redirect
from django.contrib.auth import authenticate, login, logout
from django.contrib import messages

def index(request):
    canchas = Cancha.objects.all()[:2]  # solo las primeras 2
    return render(request,'core/index.html', {'canchas': canchas})

def contacto(request):
    return render(request, 'core/contacto.html')

def promociones(request):
    return render(request, 'core/promociones.html')

def login_view(request): 
    if request.method == 'POST':
        username = request.POST['username']
        password = request.POST['password']

        user = authenticate(request, username=username, password=password)
        if user is not None:
            login(request, user)
            
            # Obtener el objeto Usuario de tu tabla Oracle
            try:
                usuario = Usuario.objects.get(email=user.email)
                request.session['usuario_id'] = usuario.id_usuario
            except Usuario.DoesNotExist:
                messages.error(request, "No se encontró tu usuario en la base de datos interna.")
                return redirect('index')

            messages.success(request, f'¡Bienvenido {user.username}! Has iniciado sesión correctamente.')


            next_url = request.GET.get('next') or request.POST.get('next') or 'index'
            return redirect(next_url)
        else:
            messages.error(request, 'Correo o contraseña incorrectos.')
            return redirect('index')

    return redirect('index')


def logout_view(request):
    logout(request)
    messages.success(request, 'Has cerrado sesión correctamente.')
    return redirect('index')




def registro(request):
    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        apellido = request.POST.get('apellido', '').strip()
        email = request.POST.get('email', '').strip()
        password = request.POST.get('password', '').strip()

        print("🟢 [DEBUG] POST recibido en /registro/")
        print(f"👉 Nombre: {nombre}")
        print(f"👉 Apellido: {apellido}")
        print(f"👉 Email: {email}")

        # Validaciones básicas
        if not nombre or not apellido or not email or not password:
            print("⚠️ Faltan campos obligatorios.")
            messages.error(request, 'Por favor completa todos los campos.')
            return redirect('registro')

        # Validar duplicados
        if User.objects.filter(username=email).exists() or Usuario.objects.filter(email=email).exists():
            print("⚠️ El correo ya existe.")
            messages.error(request, 'Ya existe una cuenta registrada con este correo.')
            return redirect('registro')

        try:
            with transaction.atomic():
                print("🚀 Creando usuario en Django auth_user...")
                user = User.objects.create_user(
                    username=email,
                    email=email,
                    password=password,
                    first_name=nombre,
                    last_name=apellido
                )
                user.save()
                print("✅ auth_user creado.")

                # Generar ID manual
                next_id = 1
                ultimo = Usuario.objects.all().order_by('-id_usuario').first()
                if ultimo:
                    next_id = ultimo.id_usuario + 1

                print(f"🧮 Próximo ID asignado: {next_id}")

                Usuario.objects.create(
                    id_usuario=next_id,
                    nombre=nombre,
                    apellido=apellido,
                    email=email,
                    password=make_password(password)
                )
                print("✅ Registro insertado correctamente en tabla USUARIO.")

            messages.success(request, 'Cuenta creada exitosamente. ¡Ya puedes iniciar sesión!')
            print("🎉 [ÉXITO] Usuario creado en ambas tablas correctamente.")
            return redirect('index')

        except Exception as e:
            print("❌ [ERROR] Falló el registro del usuario.")
            print(f"📄 Detalle: {str(e)}")
            messages.error(request, f'Ocurrió un error al registrar: {str(e)}')
            return redirect('registro')

    else:
        print("ℹ️ Carga inicial del formulario de registro.")
    return render(request, 'core/registro.html')

def olvide_contrasena(request):
    if request.method == 'POST':
        email = request.POST.get('email', '').strip()
        print(f"📩 [DEBUG] Solicitud de recuperación recibida para: {email}")

        # Validar si existe el correo
        if not email:
            print("⚠️ [WARN] No se ingresó correo.")
            messages.error(request, "Por favor ingresa tu correo electrónico.")
            return redirect('index')

        try:
            user = User.objects.get(email=email)
        except User.DoesNotExist:
            print("❌ [ERROR] No existe usuario con ese correo.")
            messages.error(request, "No existe una cuenta registrada con ese correo.")
            return redirect('index')

        # Generar token seguro y UID
        token = default_token_generator.make_token(user)
        uid = urlsafe_base64_encode(force_bytes(user.pk))
        link = request.build_absolute_uri(f"/restablecer/{uid}/{token}/")

        # Enviar correo
        asunto = "Restablece tu contraseña - CanchaYa"
        mensaje = f"""
Hola {user.first_name or 'usuario'},

Has solicitado restablecer tu contraseña en CanchaYa ⚽.
Haz clic en el siguiente enlace para crear una nueva contraseña:

{link}

Si no solicitaste esto, ignora este mensaje.

Atentamente,
El equipo de CanchaYa 💚
"""

        try:
            send_mail(
                asunto,
                mensaje,
                'canchasya.duoc@gmail.com',  # tu cuenta emisora
                [email],
                fail_silently=False,
            )
            print(f"✅ [INFO] Enlace de recuperación enviado a {email}")
            messages.success(request, "Te enviamos un enlace para restablecer tu contraseña.")
        except Exception as e:
            print(f"❌ [ERROR] No se pudo enviar el correo: {e}")
            messages.error(request, "Hubo un problema al enviar el correo. Intenta nuevamente más tarde.")

        return redirect('index')

    return redirect('index')

def restablecer_contrasena(request, uidb64, token):
    try:
        uid = force_str(urlsafe_base64_decode(uidb64))
        user = User.objects.get(pk=uid)
    except (TypeError, ValueError, OverflowError, User.DoesNotExist):
        print("❌ [ERROR] Token o usuario inválido.")
        user = None

    if user is not None and default_token_generator.check_token(user, token):
        print(f"🟢 [INFO] Enlace válido para {user.email}")

        # 🔹 Si el usuario envía nueva contraseña
        if request.method == 'POST':
            nueva = request.POST.get('password')
            confirmar = request.POST.get('confirmar')

            if not nueva or not confirmar:
                messages.error(request, "Debes ingresar ambas contraseñas.")
                return redirect(request.path)

            if nueva != confirmar:
                messages.error(request, "Las contraseñas no coinciden.")
                return redirect(request.path)

            try:
                with transaction.atomic():
                    # Actualizar contraseña en Django (auth_user)
                    user.set_password(nueva)
                    user.save()

                    # Actualizar también en tabla Oracle 'usuario'
                    Usuario.objects.filter(email=user.email).update(password=make_password(nueva))
                    print(f"🔒 [OK] Contraseña actualizada para {user.email}")

                messages.success(request, "Tu contraseña fue restablecida exitosamente.")
                return redirect('index')

            except Exception as e:
                print(f"❌ [ERROR] Al actualizar contraseña: {e}")
                messages.error(request, "Ocurrió un error al actualizar tu contraseña.")
                return redirect('index')

        # 🔹 Si solo abre el enlace (GET)
        print("📩 [INFO] Mostrando modal de restablecer contraseña...")
        return render(request, 'core/index.html', {'abrir_reset': True, 'uid': uidb64, 'token': token})

    else:
        print("❌ [ERROR] Enlace expirado o inválido.")
        messages.error(request, "El enlace de recuperación no es válido o ha expirado.")
        return redirect('index')
    

# views.py
def canchas(request):
    canchas_list = Cancha.objects.filter(estado=True)

    tipo = request.GET.get('tipo')
    precio = request.GET.get('precio')
    hora = request.GET.get('hora')
    fecha_str = request.GET.get('fecha')

    if tipo:
        canchas_list = canchas_list.filter(tipo_cancha__nombre__icontains=tipo)

    if precio:
        try:
            precio_max = int(precio)
            canchas_list = canchas_list.filter(precio__lte=precio_max)
        except ValueError:
            pass

    if hora:
        try:
            h, m = map(int, hora.split(':'))
            hora_obj = time(h, m)
            canchas_list = canchas_list.filter(hora_inicio__lte=hora_obj, hora_fin__gte=hora_obj)
        except ValueError:
            pass

    if fecha_str:
        try:
            fecha = datetime.strptime(fecha_str, "%Y-%m-%d").date()
            canchas_list = canchas_list.exclude(reserva__fecha=fecha, reserva__estado='A').distinct()
        except ValueError:
            pass

    paginator = Paginator(canchas_list, 5)
    page_number = request.GET.get('page')
    canchas = paginator.get_page(page_number)

    tipos = TipoCancha.objects.all()
    horas = [f"{h:02d}:00" for h in range(8, 24)]

    # 🟢 Asegurar formato ISO para HTML5
    hoy = date.today().strftime("%Y-%m-%d")
    max_fecha = (date.today() + timedelta(days=30)).strftime("%Y-%m-%d")

    return render(request, 'core/canchas.html', {
        'canchas': canchas,
        'tipos': tipos,
        'horas': horas,
        'hoy': hoy,
        'max_fecha': max_fecha,
    })



from django.contrib.auth.decorators import login_required
from django.db import connections

@login_required
def perfil(request):
    email = request.user.email
    nombre = None
    apellido = None

    try:
        with connections['default'].cursor() as cursor:
            cursor.execute("""
                SELECT NOMBRE, APELLIDO
                FROM USUARIO
                WHERE EMAIL = :email
            """, {'email': email})
            row = cursor.fetchone()
            if row:
                nombre, apellido = row
    except Exception as e:
        print(f"⚠️ [Perfil] Error obteniendo datos de Oracle: {e}")

    context = {
        'nombre': nombre,
        'apellido_usuario': apellido,
        'email_usuario': email,
    }
    return render(request, 'core/perfil.html', context)


@login_required
def editar_perfil(request):
    """
    Permite actualizar nombre y apellido del usuario logueado en Oracle.
    """
    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        apellido = request.POST.get('apellido', '').strip()
        email = request.user.email

        # 🧩 Validaciones backend
        if len(nombre) < 3 or len(apellido) < 3:
            return JsonResponse({'ok': False, 'msg': 'El nombre y apellido deben tener al menos 3 caracteres.'})

        if not re.match(r'^[A-Za-zÁÉÍÓÚáéíóúÑñ ]+$', nombre) or not re.match(r'^[A-Za-zÁÉÍÓÚáéíóúÑñ ]+$', apellido):
            return JsonResponse({'ok': False, 'msg': 'Solo se permiten letras en los campos.'})

        try:
            with connections['default'].cursor() as cursor:
                cursor.execute("""
                    UPDATE USUARIO
                    SET NOMBRE = :nombre, APELLIDO = :apellido
                    WHERE EMAIL = :email
                """, {'nombre': nombre, 'apellido': apellido, 'email': email})
            
            # También sincroniza con el modelo auth_user de Django
            request.user.first_name = nombre
            request.user.last_name = apellido
            request.user.save()

            return JsonResponse({'ok': True, 'msg': 'Datos actualizados correctamente.'})
        except Exception as e:
            print(f"❌ [ERROR] Falló la actualización del perfil: {e}")
            return JsonResponse({'ok': False, 'msg': 'Ocurrió un error al actualizar los datos.'})

    return JsonResponse({'ok': False, 'msg': 'Método inválido.'})



def api_stock_equipamientos(request):
    fecha_str = request.GET.get('fecha')
    horario_id = request.GET.get('horario_id')

    if not fecha_str or not horario_id:
        return JsonResponse({'error': 'Faltan parámetros (fecha o horario_id).'}, status=400)

    try:
        fecha = datetime.strptime(fecha_str, "%Y-%m-%d").date()
    except ValueError:
        return JsonResponse({'error': 'Formato de fecha inválido (usar YYYY-MM-DD).'}, status=400)

    try:
        horario_id = int(horario_id)  # 🔹 convierte a entero para que el filtro funcione
    except ValueError:
        return JsonResponse({'error': 'ID de horario inválido.'}, status=400)

    equipamientos = Equipamiento.objects.all()
    data = []

    for eq in equipamientos:
        reservado = ReservaEquipamiento.objects.filter(
            reserva__fecha=fecha,
            reserva__horario_id=horario_id,
            reserva__estado='A',
            equipamiento=eq
        ).aggregate(total_reservado=Sum('cantidad'))['total_reservado'] or 0

        disponible = max(eq.stock - reservado, 0)

        data.append({
            'id': eq.id_equipamiento,
            'nombre': eq.nombre,
            'stock_total': eq.stock,
            'reservado': reservado,
            'disponible': disponible,
            'precio': eq.precio,
        })

    return JsonResponse({'equipamientos': data})


def reserva(request, id_cancha):
    cancha = get_object_or_404(Cancha, pk=id_cancha)

    # 🟩 Equipamientos activos asociados al tipo de cancha
    equipamientos_disponibles = Equipamiento.objects.filter(
        tipos_cancha=cancha.tipo_cancha,
        estado=True
    )

    # Fechas mínimas y máximas
    hoy = date.today()
    fecha_max = hoy + timedelta(days=30)

    # Manejar selección de fecha
    fecha_str = request.GET.get('fecha')
    if fecha_str:
        try:
            fecha_seleccionada = datetime.strptime(fecha_str, "%Y-%m-%d").date()
        except ValueError:
            fecha_seleccionada = hoy
    else:
        fecha_seleccionada = hoy

    # 🟩 Horarios activos dentro del rango de la cancha
    horarios_qs = Horario.objects.filter(
        hora_inicio__gte=cancha.hora_inicio,
        hora_fin__lte=cancha.hora_fin,
        estado=True
    ).order_by('hora_inicio')

    # 🟢 Incluir el precio real (según Tarifa activa o el precio base)
    horarios_disponibles = []
    for h in horarios_qs:
        tarifa = Tarifa.objects.filter(
            cancha=cancha,
            horario=h,
            estado=True
        ).first()
        precio_final = tarifa.precio if tarifa else cancha.precio

        horarios_disponibles.append({
            "id_horario": h.id_horario,
            "hora_inicio": h.hora_inicio.strftime("%H:%M"),
            "hora_fin": h.hora_fin.strftime("%H:%M"),
            "precio": precio_final,
        })

    # 🟩 Horarios ocupados para la fecha seleccionada (solo reservas activas)
    reservas_ocupadas = Reserva.objects.filter(
        cancha=cancha,
        fecha=fecha_seleccionada,
        estado='A'
    ).values_list('horario_id', flat=True)

    horarios_ocupados = list(reservas_ocupadas)

    # 🟩 Promociones activas
    promociones_activas = Promocion.objects.filter(activo=True)

    context = {
        'cancha': cancha,
        'horarios_disponibles': horarios_disponibles,
        'horarios_ocupados': horarios_ocupados,
        'fecha_actual': hoy,
        'fecha_max': fecha_max,
        'fecha_seleccionada': fecha_seleccionada,
        'equipamientos_disponibles': equipamientos_disponibles,
        'promociones': promociones_activas,
    }

    return render(request, 'core/reserva.html', context)




def comprobante(request):
    return render(request, 'core/comprobante.html')

def cuenta(request):
    return render(request, 'core/cuenta.html')

def modificarCuenta(request):
    return render(request, 'core/modificarCuenta.html')


def historialReserva(request):
    usuario_id = request.session.get('usuario_id')
    if not usuario_id:
        messages.error(request, "Debes iniciar sesión para ver tus reservas.")
        return redirect('login')

    usuario = get_object_or_404(Usuario, id_usuario=usuario_id)
    
    reservas_list = (
        Reserva.objects
        .filter(usuario=usuario)
        .select_related('cancha', 'horario', 'promocion')
        .prefetch_related('reservaequipamiento_set__equipamiento')
        .order_by('-fecha', '-horario__hora_inicio')
    )

    # 🔹 Paginación igual que en canchas (5 por página)
    paginator = Paginator(reservas_list, 5)
    page_number = request.GET.get('page')
    reservas = paginator.get_page(page_number)

    return render(request, 'core/historialReserva.html', {'reservas': reservas})

def detalleReserva(request, id_reserva):
    usuario_id = request.session.get('usuario_id')
    if not usuario_id:
        messages.error(request, "Debes iniciar sesión para ver el detalle de tu reserva.")
        return redirect('login')

    reserva = get_object_or_404(
        Reserva.objects
        .select_related('cancha', 'horario', 'promocion')
        .prefetch_related('reservaequipamiento_set__equipamiento'),
        id_reserva=id_reserva,
        usuario_id=usuario_id
    )

    return render(request, 'core/detalleReserva.html', {'reserva': reserva})

@login_required(login_url='login')
def cancelar_reserva(request, id_reserva):
    usuario_id = request.session.get('usuario_id')
    if not usuario_id:
        messages.error(request, "Debes iniciar sesión para cancelar una reserva.")
        return redirect('login')

    # Solo puede cancelar sus propias reservas
    reserva = get_object_or_404(
        Reserva,
        id_reserva=id_reserva,
        usuario_id=usuario_id
    )

    # 👇 Aquí cancelamos directo (sin pantalla intermedia)
    reserva.estado = 'C'  # Ajusta el valor si usas otro código para cancelada
    reserva.save()

    # 💌 Enviar correo de cancelación
    enviar_correo_reserva_cancelada(reserva)

    messages.success(request, "Tu reserva fue cancelada correctamente.")
    return redirect('historialReserva')

def validar_promocion(request, codigo):
    try:
        promo = Promocion.objects.get(codigo__iexact=codigo, activo=True)
        return JsonResponse({
            'id': promo.id_promocion,
            'codigo': promo.codigo,
            'descuento_porcentaje': promo.descuento_porcentaje,
            'descuento_fijo': float(promo.descuento_fijo),
        })
    except Promocion.DoesNotExist:
        return JsonResponse({'error': 'Código de promoción inválido'}, status=404)
    
def api_horarios_ocupados(request):
    cancha_id = request.GET.get('cancha_id')
    fecha_str = request.GET.get('fecha')

    # 🟥 Validar parámetros obligatorios
    if not cancha_id or not fecha_str:
        return JsonResponse({'error': 'Faltan parámetros (cancha_id o fecha).'}, status=400)

    # 🟩 Validar formato de fecha
    try:
        fecha = datetime.strptime(fecha_str, "%Y-%m-%d").date()
    except ValueError:
        return JsonResponse({'error': 'Formato de fecha inválido (usar YYYY-MM-DD).'}, status=400)

    # 🟨 Verificar existencia de cancha
    try:
        cancha = Cancha.objects.get(id_cancha=cancha_id)
    except Cancha.DoesNotExist:
        return JsonResponse({'error': 'La cancha no existe.'}, status=404)

    # 🟩 Obtener horarios dentro del rango de la cancha
    horarios_qs = Horario.objects.filter(
        hora_inicio__gte=cancha.hora_inicio,
        hora_fin__lte=cancha.hora_fin
    ).order_by('hora_inicio')

    # 🟩 Construir lista de horarios con precio (tarifa o base)
    horarios_disponibles = []
    for h in horarios_qs:
        tarifa = Tarifa.objects.filter(cancha=cancha, horario=h).first()
        precio_final = tarifa.precio if tarifa else cancha.precio

        horarios_disponibles.append({
            "id_horario": h.id_horario,
            "hora_inicio": h.hora_inicio.strftime("%H:%M"),
            "hora_fin": h.hora_fin.strftime("%H:%M"),
            "precio": precio_final  # 🔹 nuevo campo agregado
        })

    # 🟩 Obtener horarios ocupados (solo IDs)
    horarios_ocupados = list(
        Reserva.objects.filter(
            cancha=cancha,
            fecha=fecha,
            estado='A'
        ).values_list('horario_id', flat=True)
    )

    # 🟩 Retornar respuesta JSON
    return JsonResponse({
        'horarios_disponibles': horarios_disponibles,
        'horarios_ocupados': horarios_ocupados
    })


stripe.api_key = settings.STRIPE_SECRET_KEY
from decimal import Decimal

def crear_checkout(request):
    if request.method != "POST":
        return redirect('index')

    usuario_id = request.session.get('usuario_id')
    if not usuario_id:
        messages.error(request, "Debes iniciar sesión para reservar.")
        return redirect('login')

    cancha_id = request.POST.get('cancha_id')
    fecha = request.POST.get('fecha')
    horario_id = request.POST.get('horario_id')
    equipamientos_json = request.POST.get('equipamientos', '{}')
    codigo_promocion = request.POST.get('codigo_promocion', '').strip()

    # Validaciones básicas mínimas
    if not (cancha_id and fecha and horario_id):
        messages.error(request, "Faltan datos de la reserva.")
        return redirect('index')

    cancha = get_object_or_404(Cancha, id_cancha=cancha_id)
    horario = get_object_or_404(Horario, id_horario=horario_id)

    # -------- Recalcular montos en backend (tarifa por horario + equipamientos + promo) --------
    # Precio base según tarifa si existe; si no, precio de la cancha
    tarifa = Tarifa.objects.filter(cancha=cancha, horario=horario).first()
    precio_base = tarifa.precio if tarifa else cancha.precio

    # Equipamientos
    subtotal = int(precio_base)
    try:
        equipamientos = json.loads(equipamientos_json) if equipamientos_json else {}
    except json.JSONDecodeError:
        equipamientos = {}

    for equip_id, cantidad in equipamientos.items():
        try:
            cantidad_int = int(cantidad or 0)
        except (TypeError, ValueError):
            cantidad_int = 0
        if cantidad_int > 0:
            equip = get_object_or_404(Equipamiento, id_equipamiento=equip_id)
            subtotal += int(equip.precio) * cantidad_int

    # Promoción
    descuento = 0
    promo = None
    if codigo_promocion:
        try:
            promo = Promocion.objects.get(codigo__iexact=codigo_promocion, activo=True)
            if promo.descuento_porcentaje:
                descuento = int(subtotal * promo.descuento_porcentaje / 100)
            elif promo.descuento_fijo:
                descuento = int(promo.descuento_fijo)
            if descuento > subtotal:
                descuento = subtotal
        except Promocion.DoesNotExist:
            descuento = 0
            promo = None

    total = subtotal - descuento

    if total <= 0:
        messages.error(request, "El total debe ser mayor que 0 para continuar con el pago.")
        return redirect('reserva', id_cancha=cancha_id)

    # Guardar datos (ya con montos recalculados en backend)
    request.session['reserva_temp'] = {
        'cancha_id': cancha_id,
        'fecha': fecha,
        'horario_id': horario_id,
        'equipamientos': json.dumps(equipamientos),
        'codigo_promocion': codigo_promocion,
        'subtotal': subtotal,
        'descuento': descuento,
        'total': total
    }

    # En CLP no hay decimales
    unit_amount = int(total)

    try:
        session = stripe.checkout.Session.create(
            payment_method_types=['card'],
            line_items=[{
                'price_data': {
                    'currency': 'clp',
                    'product_data': {
                        'name': f'Reserva {cancha.nombre}',
                    },
                    'unit_amount': unit_amount,
                },
                'quantity': 1,
            }],
            mode='payment',
            success_url='http://127.0.0.1:8000/pago_exitoso/',
            cancel_url='http://127.0.0.1:8000/pago_fallido/',
        )
        return redirect(session.url, code=303)

    except Exception as e:
        messages.error(request, f"Ocurrió un error al iniciar el pago: {e}")
        return redirect('reserva', id_cancha=cancha_id)


from django.db import transaction

from django.db import transaction
from datetime import datetime, date
from django.core.mail import send_mail

def pago_exitoso(request):
    reserva_temp = request.session.get('reserva_temp')
    if not reserva_temp:
        messages.warning(request, "No se encontró información de la reserva pagada.")
        return redirect('index')

    usuario_id = request.session.get('usuario_id')
    if not usuario_id:
        messages.error(request, "Debes iniciar sesión para ver tu reserva.")
        return redirect('login')

    try:
        cancha_id = reserva_temp.get('cancha_id')
        fecha_str = reserva_temp.get('fecha')
        horario_id = reserva_temp.get('horario_id')
        codigo_promocion = reserva_temp.get('codigo_promocion', '').strip()
        equipamientos_json = reserva_temp.get('equipamientos', '{}')

        # Convertir string a date
        try:
            fecha = datetime.strptime(fecha_str, "%Y-%m-%d").date()
        except (TypeError, ValueError):
            fecha = date.today()

        cancha = get_object_or_404(Cancha, id_cancha=cancha_id)
        usuario = get_object_or_404(Usuario, id_usuario=usuario_id)
        horario = get_object_or_404(Horario, id_horario=horario_id)

        # Evitar doble reserva
        if Reserva.objects.filter(cancha=cancha, fecha=fecha, horario=horario, estado='A').exists():
            del request.session['reserva_temp']
            messages.error(request, "Este horario se reservó mientras completabas el pago.")
            return redirect('reserva', id_cancha=cancha_id)

        # Recalcular montos con tarifa por horario
        tarifa = Tarifa.objects.filter(cancha=cancha, horario=horario).first()
        precio_base = tarifa.precio if tarifa else cancha.precio
        subtotal = int(precio_base)

        try:
            equipamientos = json.loads(equipamientos_json) if equipamientos_json else {}
        except json.JSONDecodeError:
            equipamientos = {}

        for equip_id, cantidad in equipamientos.items():
            try:
                cantidad_int = int(cantidad or 0)
            except (TypeError, ValueError):
                cantidad_int = 0
            if cantidad_int > 0:
                equip = get_object_or_404(Equipamiento, id_equipamiento=equip_id)
                subtotal += int(equip.precio) * cantidad_int

        descuento = 0
        promo = None
        if codigo_promocion:
            try:
                promo = Promocion.objects.get(codigo__iexact=codigo_promocion, activo=True)
                if promo.descuento_porcentaje:
                    descuento = int(subtotal * promo.descuento_porcentaje / 100)
                elif promo.descuento_fijo:
                    descuento = int(promo.descuento_fijo)
                if descuento > subtotal:
                    descuento = subtotal
            except Promocion.DoesNotExist:
                descuento = 0
                promo = None

        total = subtotal - descuento

        # Crear la reserva
        with transaction.atomic():
            reserva = Reserva.objects.create(
                fecha=fecha,
                subtotal=subtotal,
                descuento=descuento,
                total=total,
                estado='A',
                cancha=cancha,
                usuario=usuario,
                horario=horario,
                promocion=promo
            )

            for equip_id, cantidad in equipamientos.items():
                try:
                    cantidad_int = int(cantidad or 0)
                except (TypeError, ValueError):
                    cantidad_int = 0
                if cantidad_int > 0:
                    equip = get_object_or_404(Equipamiento, id_equipamiento=equip_id)
                    ReservaEquipamiento.objects.create(
                        reserva=reserva,
                        equipamiento=equip,
                        cantidad=cantidad_int
                    )

        # Limpiar la sesión temporal
        del request.session['reserva_temp']

        # Enviar correo de confirmación
        try:
            hora_ini = horario.hora_inicio.strftime('%H:%M')
            hora_fin = horario.hora_fin.strftime('%H:%M')
            asunto = "Confirmación de reserva - CanchaYa"
            mensaje = f"""
Hola {usuario.nombre or 'jugador'},

Tu reserva ha sido confirmada con éxito ⚽

📅 Fecha: {fecha.strftime('%d/%m/%Y')}
🕒 Horario: {hora_ini} - {hora_fin}
📍 Cancha: {cancha.nombre}
🏠 Dirección: {cancha.direccion}
💰 Total pagado: ${total}

¡Gracias por preferir CanchaYa! Nos vemos en la cancha 💚
"""
            send_mail(
                asunto,
                mensaje,
                'canchasya.duoc@gmail.com',  # emisor
                [usuario.email],             # tu modelo Usuario usa 'email'
                fail_silently=False,
            )
            print(f"✅ [INFO] Correo de confirmación enviado a {usuario.email}")
        except Exception as e:
            print(f"❌ [ERROR] No se pudo enviar el correo de confirmación: {e}")

        # Render final
        return render(request, 'core/pago_exitoso.html', {
            'reserva': reserva,
            'fecha': fecha
        })

    except Exception as e:
        messages.error(request, f"Ocurrió un error al registrar la reserva: {str(e)}")
        return redirect('index')


def pago_fallido(request):
    return render(request, 'core/pago_fallido.html')


import requests

def obtener_clima(request):
    lat = request.GET.get('lat')
    lon = request.GET.get('lon')
    fecha = request.GET.get('fecha')

    if not lat or not lon:
        return JsonResponse({'error': 'Faltan coordenadas'}, status=400)

    API_KEY = '798b4f4d195871c9ca50aad795bd6420'
    url = f"https://api.openweathermap.org/data/2.5/forecast?lat={lat}&lon={lon}&appid={API_KEY}&units=metric&lang=es"

    try:
        r = requests.get(url, timeout=5)
        data = r.json()
        if 'list' not in data:
            return JsonResponse({'error': 'Respuesta inválida de OpenWeather'}, status=500)

        fecha_reserva = datetime.strptime(fecha, "%Y-%m-%d").date()
        # Fecha máxima disponible en el pronóstico (último elemento)
        ultima_fecha = datetime.fromtimestamp(data['list'][-1]['dt']).date()

        # 🟡 Si la fecha excede el rango de 5 días, devolvemos un mensaje
        if fecha_reserva > ultima_fecha:
            return JsonResponse({
                'advertencia': 'Nuestro pronóstico está disponible solo dentro de los próximos 5 días.'
            })

        # Buscar el pronóstico más cercano al mediodía del día elegido
        pronostico = None
        for entry in data['list']:
            dt = datetime.fromtimestamp(entry['dt'])
            if dt.date() == fecha_reserva and dt.hour in [11, 12, 13]:
                pronostico = entry
                break

        # Si no se encontró exactamente al mediodía, tomamos el más cercano
        if not pronostico:
            pronostico = min(
                data['list'],
                key=lambda e: abs(datetime.fromtimestamp(e['dt']).date() - fecha_reserva)
            )

        clima = {
            'temperatura': pronostico['main']['temp'],
            'descripcion': pronostico['weather'][0]['description'].capitalize(),
            'icono': pronostico['weather'][0]['icon'],
        }
        return JsonResponse(clima)

    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)
    
def reportes_ingresos(request):
    # Parámetros de filtro (GET)
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')
    cancha_id = request.GET.get('cancha')

    # Query base: solo reservas activas
    reservas = Reserva.objects.filter(estado='A')

    if fecha_inicio:
        reservas = reservas.filter(fecha__gte=fecha_inicio)
    if fecha_fin:
        reservas = reservas.filter(fecha__lte=fecha_fin)
    if cancha_id and cancha_id != '':
        reservas = reservas.filter(cancha_id=cancha_id)

    # Agregados
    resumen = reservas.aggregate(
        total_ingresos=Sum('total'),
        total_descuentos=Sum('descuento'),
        total_subtotal=Sum('subtotal'),
        cantidad_reservas=Count('id_reserva')
    )

    # Reporte por cancha
    ingresos_por_cancha = reservas.values('cancha__nombre').annotate(
        ingresos=Sum('total'),
        reservas=Count('id_reserva')
    ).order_by('-ingresos')

    context = {
        'reservas': reservas,
        'resumen': resumen,
        'ingresos_por_cancha': ingresos_por_cancha,
        'canchas': Cancha.objects.all(),
        'fecha_inicio': fecha_inicio,
        'fecha_fin': fecha_fin,
        'cancha_id': cancha_id,
    }
    return render(request, 'core/reportes_ingresos.html', context)


# 🟢 EXPORTAR A EXCEL
def exportar_ingresos_excel(request):
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')
    cancha_id = request.GET.get('cancha')

    reservas = Reserva.objects.filter(estado='A')

    # ✅ Validaciones seguras
    if fecha_inicio and fecha_inicio.lower() != "none":
        reservas = reservas.filter(fecha__gte=fecha_inicio)
    if fecha_fin and fecha_fin.lower() != "none":
        reservas = reservas.filter(fecha__lte=fecha_fin)
    if cancha_id and cancha_id.lower() != "none" and cancha_id != "":
        reservas = reservas.filter(cancha_id=cancha_id)

    # 🧾 Construcción de datos
    data = []
    for r in reservas:
        fecha_str = r.fecha.strftime("%Y-%m-%d") if r.fecha else None
        data.append({
            'Fecha': fecha_str,
            'Cancha': r.cancha.nombre,
            'Usuario': f"{r.usuario.nombre} {r.usuario.apellido}",
            'Subtotal': r.subtotal,
            'Descuento': r.descuento,
            'Total': r.total,
        })

    df = pd.DataFrame(data)

    # Convertir la fecha en datetime real
    if not df.empty:
        df['Fecha'] = pd.to_datetime(df['Fecha'], errors='coerce')

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Reservas')

        # 📊 Hoja resumen
        if not df.empty:
            resumen = df.groupby('Cancha').agg({
                'Subtotal': 'sum',
                'Descuento': 'sum',
                'Total': 'sum'
            }).reset_index()
            resumen.to_excel(writer, index=False, sheet_name='Resumen')

        # 🎨 Ajustes de formato visual
        workbook = writer.book
        ws = writer.sheets['Reservas']

        # ✅ Auto-filtro
        ws.auto_filter.ref = ws.dimensions

        # ✅ Ajustar ancho de columnas automáticamente
        for col in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    cell_length = len(str(cell.value))
                    if cell_length > max_length:
                        max_length = cell_length
            adjusted_width = (max_length + 2)
            ws.column_dimensions[col_letter].width = adjusted_width

        # ✅ Formato de fechas (columna “Fecha”)
        for cell in ws['A']:
            if cell.row == 1:
                continue
            cell.number_format = 'YYYY-MM-DD'

        # ✅ Formato de números con separador de miles
        for col_idx, col_name in enumerate(df.columns, 1):
            if col_name in ['Subtotal', 'Descuento', 'Total']:
                col_letter = get_column_letter(col_idx)
                for cell in ws[col_letter]:
                    if cell.row == 1:
                        continue
                    cell.number_format = '#,##0'

    output.seek(0)
    response = HttpResponse(
        output,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="reporte_ingresos.xlsx"'
    return response


# 🟥 EXPORTAR A PDF
def exportar_ingresos_pdf(request):
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')
    cancha_id = request.GET.get('cancha')

    reservas = Reserva.objects.filter(estado='A')

    # ✅ Validaciones seguras
    if fecha_inicio and fecha_inicio.lower() != "none":
        reservas = reservas.filter(fecha__gte=fecha_inicio)
    if fecha_fin and fecha_fin.lower() != "none":
        reservas = reservas.filter(fecha__lte=fecha_fin)
    if cancha_id and cancha_id.lower() != "none" and cancha_id != "":
        reservas = reservas.filter(cancha_id=cancha_id)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
    elements = []
    styles = getSampleStyleSheet()

    # 🧾 Título
    elements.append(Paragraph("Reporte de Ingresos - CanchaYa", styles['Title']))
    elements.append(Spacer(1, 12))

    # 🧮 Encabezado de tabla
    data = [['Fecha', 'Cancha', 'Usuario', 'Subtotal', 'Descuento', 'Total']]

    for r in reservas:
        data.append([
            str(r.fecha),
            r.cancha.nombre,
            f"{r.usuario.nombre} {r.usuario.apellido}",
            f"${r.subtotal:,}",
            f"${r.descuento:,}",
            f"${r.total:,}",
        ])

    # 🧱 Tabla PDF
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#198754")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey]),
    ]))

    elements.append(table)
    doc.build(elements)

    pdf = buffer.getvalue()
    buffer.close()
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="reporte_ingresos.pdf"'
    response.write(pdf)
    return response


def reportes_ocupaciones(request):
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')
    cancha_id = request.GET.get('cancha')

    reservas = Reserva.objects.filter(estado='A')

    if fecha_inicio and fecha_inicio.lower() != "none":
        reservas = reservas.filter(fecha__gte=fecha_inicio)
    if fecha_fin and fecha_fin.lower() != "none":
        reservas = reservas.filter(fecha__lte=fecha_fin)
    if cancha_id and cancha_id.lower() != "none" and cancha_id != "":
        reservas = reservas.filter(cancha_id=cancha_id)

    # 📊 Datos resumen
    resumen = {
        'total_reservas': reservas.count(),
        'canchas_ocupadas': reservas.values('cancha').distinct().count(),
        'dias_ocupados': reservas.values('fecha').distinct().count(),
    }

    # 📅 Agrupar reservas por día de la semana (compatibles con Oracle)
    reservas_por_dia = (
        reservas.annotate(dia_semana=ExtractWeekDay('fecha'))
        .values('dia_semana')
        .annotate(total=Count('id_reserva'))
        .order_by('dia_semana')
    )

    # ⏰ Agrupar reservas por hora
    reservas_por_hora = (
        reservas.values('horario__hora_inicio')
        .annotate(total=Count('id_reserva'))
        .order_by('horario__hora_inicio')
    )

    # 🗓️ Mapeo de días según Oracle (1=Domingo, 7=Sábado)
    dias_map = {
        1: 'Dom',
        2: 'Lun',
        3: 'Mar',
        4: 'Mié',
        5: 'Jue',
        6: 'Vie',
        7: 'Sáb',
    }

    for d in reservas_por_dia:
        d['dia_semana'] = dias_map.get(d['dia_semana'], '?')

    context = {
        'fecha_inicio': fecha_inicio,
        'fecha_fin': fecha_fin,
        'cancha_id': cancha_id,
        'canchas': Cancha.objects.all(),
        'reservas': reservas,
        'resumen': resumen,
        'reservas_por_dia': reservas_por_dia,
        'reservas_por_hora': reservas_por_hora,
    }

    return render(request, 'core/reportes_ocupaciones.html', context)

def centro_reportes(request):
    return render(request, 'core/centro_reportes.html')

def exportar_ocupaciones_excel(request):
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')
    cancha_id = request.GET.get('cancha')

    reservas = Reserva.objects.filter(estado='A')

    if fecha_inicio and fecha_inicio.lower() != "none":
        reservas = reservas.filter(fecha__gte=fecha_inicio)
    if fecha_fin and fecha_fin.lower() != "none":
        reservas = reservas.filter(fecha__lte=fecha_fin)
    if cancha_id and cancha_id.lower() != "none" and cancha_id != "":
        reservas = reservas.filter(cancha_id=cancha_id)

    # 📊 Agrupar datos
    reservas_por_dia = (
        reservas.annotate(dia_semana=ExtractWeekDay('fecha'))
        .values('dia_semana')
        .annotate(total=Count('id_reserva'))
        .order_by('dia_semana')
    )

    reservas_por_hora = (
        reservas.values('horario__hora_inicio')
        .annotate(total=Count('id_reserva'))
        .order_by('horario__hora_inicio')
    )

    dias_map = {
        1: 'Dom', 2: 'Lun', 3: 'Mar', 4: 'Mié',
        5: 'Jue', 6: 'Vie', 7: 'Sáb'
    }
    for d in reservas_por_dia:
        d['dia_semana'] = dias_map.get(d['dia_semana'], '?')

    # 🧾 Detalle reservas
    data = []
    for r in reservas:
        data.append({
            'Fecha': r.fecha.strftime("%Y-%m-%d") if r.fecha else "",
            'Cancha': r.cancha.nombre,
            'Horario': f"{r.horario.hora_inicio} - {r.horario.hora_fin}",
            'Usuario': f"{r.usuario.nombre} {r.usuario.apellido}",
        })

    df_detalle = pd.DataFrame(data)
    df_dias = pd.DataFrame(list(reservas_por_dia))
    df_horas = pd.DataFrame(list(reservas_por_hora))

    # Excel en memoria
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # Hoja 1: detalle
        df_detalle.to_excel(writer, index=False, sheet_name='Reservas')

        # Hoja 2: por día
        if not df_dias.empty:
            df_dias.rename(columns={'dia_semana': 'Día', 'total': 'Total Reservas'}, inplace=True)
            df_dias.to_excel(writer, index=False, sheet_name='Por Día')

        # Hoja 3: por hora
        if not df_horas.empty:
            df_horas.rename(columns={'horario__hora_inicio': 'Hora Inicio', 'total': 'Total Reservas'}, inplace=True)
            df_horas.to_excel(writer, index=False, sheet_name='Por Hora')

        workbook = writer.book
        for ws_name in workbook.sheetnames:
            ws = writer.sheets[ws_name]
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                max_length = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max_length + 2

    output.seek(0)
    response = HttpResponse(
        output,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="reporte_ocupaciones.xlsx"'
    return response


def exportar_ocupaciones_pdf(request):
    from .models import Reserva
    from django.db.models import Count
    from django.db.models.functions import ExtractWeekDay

    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')
    cancha_id = request.GET.get('cancha')

    reservas = Reserva.objects.filter(estado='A')

    if fecha_inicio and fecha_inicio.lower() != "none":
        reservas = reservas.filter(fecha__gte=fecha_inicio)
    if fecha_fin and fecha_fin.lower() != "none":
        reservas = reservas.filter(fecha__lte=fecha_fin)
    if cancha_id and cancha_id.lower() != "none" and cancha_id != "":
        reservas = reservas.filter(cancha_id=cancha_id)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
    elements = []
    styles = getSampleStyleSheet()

    # Título
    elements.append(Paragraph("Reporte de Ocupaciones - CanchaYa", styles['Title']))
    elements.append(Spacer(1, 12))

    # Tabla
    data = [['Fecha', 'Cancha', 'Horario', 'Usuario']]
    for r in reservas:
        data.append([
            str(r.fecha),
            r.cancha.nombre,
            f"{r.horario.hora_inicio} - {r.horario.hora_fin}",
            f"{r.usuario.nombre} {r.usuario.apellido}",
        ])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0d6efd")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey]),
    ]))

    elements.append(table)
    doc.build(elements)

    pdf = buffer.getvalue()
    buffer.close()
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="reporte_ocupaciones.pdf"'
    response.write(pdf)
    return response


def centroGestion(request):
    return render(request, 'core/centroGestion.html')

def gestionComercial(request):
    return render(request, 'core/gestionComercial.html')

def gestionCuentas(request):
    return render(request, 'core/gestionCuentas.html')

def crudOperadores(request):
    return render(request, 'core/crudOperadores.html')


def crudAdministradores(request):
    return render(request, 'core/crudAdministradores.html')


@login_required
def crudUsuarios(request):
    usuarios = Usuario.objects.all().order_by('id_usuario')
    return render(request, 'core/crudUsuarios.html', {'usuarios': usuarios})

@login_required
def usuario_edit(request, id_usuario):
    usuario = get_object_or_404(Usuario, id_usuario=id_usuario)
    old_email = usuario.email  # por si cambia el correo

    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        apellido = request.POST.get('apellido', '').strip()
        email = request.POST.get('email', '').strip()

        if not nombre or not apellido or not email:
            messages.error(request, 'Todos los campos son obligatorios.')
            return redirect('usuario_edit', id_usuario=id_usuario)

        with transaction.atomic():
            # actualizamos en tabla USUARIO (Oracle)
            usuario.nombre = nombre
            usuario.apellido = apellido
            usuario.email = email
            usuario.save()

            # sincronizar también con auth_user
            user = User.objects.filter(email=old_email).first()
            if user:
                user.first_name = nombre
                user.last_name = apellido
                user.email = email
                user.username = email  # porque tú usas el correo como username
                user.save()

        messages.success(request, 'Usuario modificado correctamente.')
        return redirect('crudUsuarios')

    # GET -> mostrar formulario
    return render(request, 'core/usuario_form.html', {'usuario': usuario})


@login_required
def usuario_delete(request, id_usuario):
    usuario = get_object_or_404(Usuario, id_usuario=id_usuario)
    old_email = usuario.email

    if request.method == 'POST':
        with transaction.atomic():
            # borrar auth_user si existe
            user = User.objects.filter(email=old_email).first()
            if user:
                user.delete()

            # borrar en tabla USUARIO
            usuario.delete()

        messages.success(request, 'Usuario eliminado correctamente.')
        return redirect('crudUsuarios')

    # si alguien entra por GET, lo mandamos de vuelta
    return redirect('crudUsuarios')

@login_required(login_url='login')
def crudReservas(request):
    reservas = Reserva.objects.select_related('cancha', 'usuario', 'horario', 'promocion').all().order_by('-fecha')
    return render(request, 'core/crudReservas.html', {'reservas': reservas})


@login_required(login_url='login')
def reserva_create(request):
    if request.method == 'POST':
        form = ReservaForm(request.POST)
        if form.is_valid():
            form.save()
            return JsonResponse({'success': True})
        return JsonResponse({'success': False, 'errors': form.errors})
    else:
        form = ReservaForm()
    return render(request, 'core/reserva_form.html', {'form': form})


@login_required(login_url='login')
def reserva_edit(request, pk):
    reserva = get_object_or_404(Reserva, pk=pk)
    if request.method == 'POST':
        form = ReservaForm(request.POST, instance=reserva)
        if form.is_valid():
            form.save()
            return JsonResponse({'success': True})
        return JsonResponse({'success': False, 'errors': form.errors})
    else:
        form = ReservaForm(instance=reserva)
    return render(request, 'core/reserva_form.html', {'form': form})



@login_required(login_url='login')
def reserva_delete(request, pk):
    """Cancela una reserva (eliminación lógica desde el CRUD, con notificación al usuario)"""
    reserva = get_object_or_404(Reserva, pk=pk)
    
    if request.method == 'POST':
        reserva.estado = 'C'  # 'C' = Cancelada
        reserva.save()

        # 💌 Enviar correo al usuario dueño de la reserva
        try:
            enviar_correo_reserva_cancelada(reserva)
        except Exception as e:
            print(f"Error al enviar correo de cancelación: {e}")

        return JsonResponse({'success': True})
    
    return JsonResponse({'success': False})

@login_required(login_url='login')
def reserva_activate(request, pk):
    """Reactivar una reserva cancelada"""
    reserva = get_object_or_404(Reserva, pk=pk)
    reserva.estado = 'A'  # 'A' = Activa
    reserva.save()
    return redirect('crudReservas')

@login_required(login_url='login')
def crudEquipamientos(request):
    equipamientos = Equipamiento.objects.prefetch_related('tipos_cancha').all().order_by('nombre')
    return render(request, 'core/crudEquipamientos.html', {'equipamientos': equipamientos})


@login_required(login_url='login')
def equipamiento_create(request):
    if request.method == 'POST':
        form = EquipamientoForm(request.POST)
        if form.is_valid():
            form.save()
            return JsonResponse({'success': True})
        return JsonResponse({'success': False, 'errors': form.errors})
    else:
        form = EquipamientoForm()
    return render(request, 'core/equipamiento_form.html', {'form': form})


@login_required(login_url='login')
def equipamiento_edit(request, pk):
    equipamiento = get_object_or_404(Equipamiento, pk=pk)
    if request.method == 'POST':
        form = EquipamientoForm(request.POST, instance=equipamiento)
        if form.is_valid():
            form.save()
            return JsonResponse({'success': True})
        return JsonResponse({'success': False, 'errors': form.errors})
    else:
        form = EquipamientoForm(instance=equipamiento)
    return render(request, 'core/equipamiento_form.html', {'form': form})


@login_required(login_url='login')
def equipamiento_delete(request, pk):
    """Desactiva un equipamiento (eliminación lógica)"""
    equipamiento = get_object_or_404(Equipamiento, pk=pk)
    if request.method == 'POST':
        equipamiento.estado = False
        equipamiento.save()
        return JsonResponse({'success': True})
    return JsonResponse({'success': False})


@login_required(login_url='login')
def equipamiento_activate(request, pk):
    """Reactivar un equipamiento desactivado"""
    equipamiento = get_object_or_404(Equipamiento, pk=pk)
    equipamiento.estado = True
    equipamiento.save()
    return redirect('crudEquipamientos')


@login_required(login_url='login')
def crudTarifas(request):
    tarifas = Tarifa.objects.select_related('cancha', 'horario').all()
    return render(request, 'core/crudTarifas.html', {'tarifas': tarifas})


@login_required
def tarifa_create(request):
    """
    Crear una nueva tarifa.
    """
    if request.method == 'POST':
        form = TarifaForm(request.POST)
        if form.is_valid():
            cancha = form.cleaned_data['cancha']
            horario = form.cleaned_data['horario']

            # Evitar duplicados (unique_together)
            if Tarifa.objects.filter(cancha=cancha, horario=horario).exists():
                messages.error(request, 'Ya existe una tarifa para esa cancha y horario.')
            else:
                form.save()
                messages.success(request, 'Tarifa creada correctamente.')
                return redirect('crudTarifas')
    else:
        form = TarifaForm()

    return render(request, 'core/tarifa_form.html', {
        'form': form,
        'modo': 'Agregar',
    })


@login_required
def tarifa_edit(request, pk):
    """
    Editar una tarifa existente.
    """
    tarifa = get_object_or_404(Tarifa, id_tarifa=pk)

    if request.method == 'POST':
        form = TarifaForm(request.POST, instance=tarifa)
        if form.is_valid():
            cancha = form.cleaned_data['cancha']
            horario = form.cleaned_data['horario']

            # Evitar duplicados al editar
            if Tarifa.objects.filter(cancha=cancha, horario=horario).exclude(id_tarifa=tarifa.id_tarifa).exists():
                messages.error(request, 'Ya existe una tarifa para esa cancha y horario.')
            else:
                form.save()
                messages.success(request, 'Tarifa actualizada correctamente.')
                return redirect('crudTarifas')
    else:
        form = TarifaForm(instance=tarifa)

    return render(request, 'core/tarifa_form.html', {
        'form': form,
        'modo': 'Editar',
        'tarifa': tarifa,
    })


@login_required(login_url='login')
def tarifa_delete(request, pk):
    """Desactiva una tarifa (eliminación lógica)"""
    tarifa = get_object_or_404(Tarifa, pk=pk)
    if request.method == 'POST':
        tarifa.estado = False
        tarifa.save()
        return JsonResponse({'success': True})
    return JsonResponse({'success': False})

@login_required(login_url='login')
def tarifa_activate(request, pk):
    tarifa = get_object_or_404(Tarifa, pk=pk)
    tarifa.estado = True
    tarifa.save()
    return redirect('crudTarifas')



@login_required(login_url='login')
def crudPromociones(request):
    promociones = Promocion.objects.all()
    return render(request, 'core/crudPromociones.html', {'promociones': promociones})


@login_required(login_url='login')
def promocion_create(request):
    if request.method == 'POST':
        form = PromocionForm(request.POST)
        if form.is_valid():
            form.save()
            return JsonResponse({'success': True})
        return JsonResponse({'success': False, 'errors': form.errors})
    else:
        form = PromocionForm()
    return render(request, 'core/promocion_form.html', {'form': form})


@login_required(login_url='login')
def promocion_edit(request, pk):
    promocion = get_object_or_404(Promocion, pk=pk)
    if request.method == 'POST':
        form = PromocionForm(request.POST, instance=promocion)
        if form.is_valid():
            form.save()
            return JsonResponse({'success': True})
        return JsonResponse({'success': False, 'errors': form.errors})
    else:
        form = PromocionForm(instance=promocion)
    return render(request, 'core/promocion_form.html', {'form': form})



@login_required(login_url='login')
def promocion_delete(request, pk):
    """Desactiva una promoción (eliminación lógica)"""
    promocion = get_object_or_404(Promocion, pk=pk)
    if request.method == 'POST':
        promocion.activo = False
        promocion.save()
        return JsonResponse({'success': True})
    return JsonResponse({'success': False})


@login_required(login_url='login')
def promocion_activate(request, pk):
    """Reactivar una promoción desactivada"""
    promocion = get_object_or_404(Promocion, pk=pk)
    promocion.activo = True
    promocion.save()
    return redirect('crudPromociones')


@login_required
def crudHorarios(request):
    horarios = Horario.objects.all().order_by('hora_inicio')
    return render(request, 'core/crudHorarios.html', {'horarios': horarios})


@login_required
def horario_create(request):
    if request.method == 'POST':
        form = HorarioForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Horario creado correctamente.')
            return redirect('crudHorarios')
        messages.error(request, 'Revisa los campos del formulario.')
    else:
        form = HorarioForm()
    return render(request, 'core/horarios_form.html', {
        'form': form,
        'modo': 'Agregar',
    })


@login_required
def horario_edit(request, pk):
    horario = get_object_or_404(Horario, id_horario=pk)
    if request.method == 'POST':
        form = HorarioForm(request.POST, instance=horario)
        if form.is_valid():
            form.save()
            messages.success(request, 'Horario actualizado correctamente.')
            return redirect('crudHorarios')
        messages.error(request, 'Revisa los campos del formulario.')
    else:
        form = HorarioForm(instance=horario)
    return render(request, 'core/horarios_form.html', {
        'form': form,
        'modo': 'Editar',
        'horario': horario,
    })



@login_required
def horario_delete(request, pk):
    horario = get_object_or_404(Horario, id_horario=pk)
    if request.method == 'POST':
        horario.estado = False
        horario.save()
    return redirect('crudHorarios')


@login_required
def horario_activate(request, pk):
    horario = get_object_or_404(Horario, id_horario=pk)
    horario.estado = True
    horario.save()
    return redirect('crudHorarios')

@login_required
def crudCanchas(request):
    # Reutilizamos tu nombre de vista, pero ahora enviamos datos reales
    canchas = Cancha.objects.select_related('tipo_cancha').order_by('id_cancha')
    return render(request, 'core/crudCanchas.html', {'canchas': canchas})

@login_required
def cancha_create(request):
    if request.method == 'POST':
        form = CanchaForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            messages.success(request, 'Cancha creada correctamente.')
            return redirect('crudCanchas')
        messages.error(request, 'Revisa los campos del formulario.')
    else:
        form = CanchaForm()
    return render(request, 'core/canchas_form.html', {'form': form, 'modo': 'Agregar'})

@login_required
def cancha_edit(request, pk):
    cancha = get_object_or_404(Cancha, id_cancha=pk)
    if request.method == 'POST':
        form = CanchaForm(request.POST, request.FILES, instance=cancha)
        if form.is_valid():
            form.save()
            messages.success(request, 'Cancha actualizada correctamente.')
            return redirect('crudCanchas')
        messages.error(request, 'Revisa los campos del formulario.')
    else:
        form = CanchaForm(instance=cancha)
    return render(request, 'core/canchas_form.html', {'form': form, 'modo': 'Editar', 'cancha': cancha})

@login_required
def cancha_delete(request, pk):
    cancha = get_object_or_404(Cancha, id_cancha=pk)
    if request.method == 'POST':
        cancha.estado = False
        cancha.save()
    # Redirige al CRUD (la vista JS se encarga del SweetAlert)
    return redirect('crudCanchas')


@login_required
def cancha_activate(request, pk):
    cancha = get_object_or_404(Cancha, id_cancha=pk)
    cancha.estado = True
    cancha.save()
    # Igual, sin messages ya que SweetAlert se encarga de feedback
    return redirect('crudCanchas')



def enviar_correo_reserva_cancelada(reserva):
    usuario = reserva.usuario
    cancha = reserva.cancha
    horario = reserva.horario

    hora_ini = horario.hora_inicio.strftime('%H:%M')
    hora_fin = horario.hora_fin.strftime('%H:%M')
    fecha_str = reserva.fecha.strftime('%d/%m/%Y')

    asunto = "Tu reserva ha sido cancelada - CanchaYa"
    mensaje = f"""
Hola {usuario.nombre or 'jugador'},

Tu reserva ha sido CANCELADA ❌

📅 Fecha: {fecha_str}
🕒 Horario: {hora_ini} - {hora_fin}
📍 Cancha: {cancha.nombre}
🏠 Dirección: {cancha.direccion}

Si tú no realizaste esta cancelación, contáctanos cuanto antes.

Atentamente,
El equipo de CanchaYa 💚
"""

    try:
        send_mail(
            asunto,
            mensaje,
            'canchasya.duoc@gmail.com',
            [usuario.email],
            fail_silently=False,
        )
        print(f"✅ [MAIL] Correo de reserva cancelada enviado a {usuario.email}")
    except Exception as e:
        print(f"❌ [MAIL] Error al enviar correo de reserva cancelada: {e}")
